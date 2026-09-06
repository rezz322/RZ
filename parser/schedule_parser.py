import io
import re
import json
import logging
from datetime import datetime, date
import httpx
import docx

from config import (
    GROUP_NAME,
    SCHEDULE_FILE,
    SEMESTER_START,
    DATA_DIR
)

logger = logging.getLogger(__name__)

DAYS_ORDER = ["ПОНЕДІЛОК", "ВІВТОРОК", "СЕРЕДА", "ЧЕТВЕР", "П’ЯТНИЦЯ"]

DAY_MAPPING = {
    0: "ПОНЕДІЛОК",
    1: "ВІВТОРОК",
    2: "СЕРЕДА",
    3: "ЧЕТВЕР",
    4: "П’ЯТНИЦЯ",
    5: "СУБОТА",
    6: "НЕДІЛЯ"
}

def get_academic_week_info(target_date: date = None) -> dict:
    """
    Calculates the academic week number (1..15+) and parity (even/odd).
    Semester starts on SEMESTER_START (Monday of week 1).
    """
    if target_date is None:
        target_date = date.today()
        
    start_dt = datetime.strptime(SEMESTER_START, "%Y-%m-%d").date()
    days_diff = (target_date - start_dt).days
    
    if days_diff < 0:
        week_num = 1
    else:
        week_num = (days_diff // 7) + 1
        
    is_even = (week_num % 2 == 0)
    parity = "even" if is_even else "odd"
    parity_ua = "Парний (знаменник)" if is_even else "Непарний (чисельник)"
    
    return {
        "week_number": week_num,
        "is_even": is_even,
        "parity": parity,
        "parity_ua": parity_ua,
        "target_date": target_date.strftime("%Y-%m-%d"),
        "day_name": DAY_MAPPING.get(target_date.weekday(), "")
    }

def clean_day_text(text: str) -> str:
    cleaned = text.replace('\n', '').replace(' ', '').upper()
    for day in DAYS_ORDER:
        day_normalized = day.replace("’", "'").replace(" ", "")
        if day_normalized in cleaned.replace("’", "'"):
            return day
    return cleaned

def parse_lesson_entry(text: str) -> list[dict]:
    """
    Parses cell text into a list of structured lesson objects with week ranges and links.
    """
    text = text.strip()
    if not text:
        return []
        
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    if not lines:
        return []
        
    week_regex = re.compile(
        r'^(\d+)-(\d+)\s*(н/?пар\.?|непарн\w*|пар\.?|парн\w*)?(.*)',
        re.IGNORECASE
    )
    
    # Extract links, meeting IDs, passwords
    def extract_metadata(body_lines):
        full_text = "\n".join(body_lines)
        
        # Zoom / Meet / Classroom urls
        urls = re.findall(r'https?://[^\s<>"]+|www\.[^\s<>"]+', full_text)
        
        # Meeting ID & Passcode
        zoom_id_match = re.search(r'(?:ідентифікатор|конференція|№:?|id:?)\s*[:.]?\s*(\d{3}[\s-]?\d{3,4}[\s-]?\d{3,4})', full_text, re.IGNORECASE)
        zoom_pwd_match = re.search(r'(?:код(?: доступа| конференції)?|пароль|pwd:?)\s*[:.]?\s*([A-Za-z0-9]+)', full_text, re.IGNORECASE)
        
        zoom_id = zoom_id_match.group(1).strip() if zoom_id_match else None
        zoom_pwd = zoom_pwd_match.group(1).strip() if zoom_pwd_match else None
        
        # Clean subject and teacher from text
        clean_lines = []
        for line in body_lines:
            # skip url lines
            if any(u in line for u in urls):
                # if line has other text before url, keep it
                for u in urls:
                    line = line.replace(u, '')
            # skip pure ID/pass lines
            if re.search(r'(ідентифікатор|код конференції|код доступа)', line, re.IGNORECASE):
                continue
            line = line.strip()
            if line:
                clean_lines.append(line)
                
        subject = clean_lines[0] if clean_lines else "Дисципліна не вказана"
        teacher = clean_lines[1] if len(clean_lines) > 1 else ""
        additional = "\n".join(clean_lines[2:]) if len(clean_lines) > 2 else ""
        
        return {
            "subject": subject,
            "teacher": teacher,
            "additional": additional,
            "urls": urls,
            "meeting_id": zoom_id,
            "passcode": zoom_pwd,
            "raw_text": full_text
        }

    results = []
    first_match = week_regex.match(lines[0])
    
    if first_match:
        # Check if line 1 also has a week spec
        if len(lines) > 1 and week_regex.match(lines[1]):
            # Two alternating lesson types (e.g. odd practice, even lab)
            m1 = first_match
            m2 = week_regex.match(lines[1])
            common_body = lines[2:]
            meta = extract_metadata(common_body)
            
            p1 = "odd" if m1.group(3) and "н" in m1.group(3).lower() else ("even" if m1.group(3) and "пар" in m1.group(3).lower() else "all")
            p2 = "odd" if m2.group(3) and "н" in m2.group(3).lower() else ("even" if m2.group(3) and "пар" in m2.group(3).lower() else "all")
            
            results.append({
                "start_week": int(m1.group(1)),
                "end_week": int(m1.group(2)),
                "parity": p1,
                "type": (m1.group(3) or "").strip() + " " + (m1.group(4) or "").strip(),
                **meta
            })
            results.append({
                "start_week": int(m2.group(1)),
                "end_week": int(m2.group(2)),
                "parity": p2,
                "type": (m2.group(3) or "").strip() + " " + (m2.group(4) or "").strip(),
                **meta
            })
        else:
            m = first_match
            p = "odd" if m.group(3) and "н" in m.group(3).lower() else ("even" if m.group(3) and "пар" in m.group(3).lower() else "all")
            meta = extract_metadata(lines[1:])
            results.append({
                "start_week": int(m.group(1)),
                "end_week": int(m.group(2)),
                "parity": p,
                "type": (m.group(3) or "").strip() + " " + (m.group(4) or "").strip(),
                **meta
            })
    else:
        # No week prefix (e.g. Фізичне виховання)
        meta = extract_metadata(lines)
        results.append({
            "start_week": 1,
            "end_week": 15,
            "parity": "all",
            "type": "",
            **meta
        })
        
    return results

async def download_and_parse_schedule(file_id: str, group_name: str = GROUP_NAME) -> dict:
    """
    Downloads the DOCX version of the schedule from Google Docs and parses it.
    Returns structured schedule dict.
    """
    export_url = f"https://docs.google.com/document/d/{file_id}/export?format=docx"
    logger.info(f"Downloading schedule DOCX from {export_url}...")
    
    async with httpx.AsyncClient(timeout=30.0, follow_redirects=True) as client:
        resp = await client.get(export_url)
        if resp.status_code != 200:
            raise RuntimeError(f"Failed to export schedule DOCX: HTTP {resp.status_code}")
        content = resp.content

    # Save backup docx
    backup_path = DATA_DIR / "schedule_latest.docx"
    with open(backup_path, "wb") as f:
        f.write(content)
        
    doc = docx.Document(io.BytesIO(content))
    if not doc.tables:
        raise ValueError("No tables found in schedule document!")
        
    table = doc.tables[0]
    
    # Identify column for target group
    header_cells = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells]
    target_col = -1
    for idx, cell_text in enumerate(header_cells):
        clean_name = cell_text.replace(" ", "")
        if group_name.replace(" ", "").lower() in clean_name.lower():
            target_col = idx
            break
            
    if target_col == -1:
        raise ValueError(f"Group '{group_name}' not found in header row: {header_cells}")
        
    logger.info(f"Found column {target_col} for group '{group_name}'")
    
    schedule_data = {
        "group": group_name,
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "file_id": file_id,
        "days": {day: {} for day in DAYS_ORDER}
    }
    
    current_day = ""
    for row_idx in range(1, len(table.rows)):
        row = table.rows[row_idx]
        raw_day = row.cells[0].text.strip()
        if raw_day:
            day_candidate = clean_day_text(raw_day)
            if day_candidate in DAYS_ORDER:
                current_day = day_candidate
                
        if not current_day:
            continue
            
        pair_num = row.cells[1].text.strip()
        if not pair_num or not pair_num.isdigit():
            continue
            
        cell_content = row.cells[target_col].text.strip()
        if not cell_content:
            continue
            
        lessons = parse_lesson_entry(cell_content)
        if lessons:
            if pair_num not in schedule_data["days"][current_day]:
                schedule_data["days"][current_day][pair_num] = []
            
            # Avoid exact duplicates
            for l in lessons:
                if l not in schedule_data["days"][current_day][pair_num]:
                    schedule_data["days"][current_day][pair_num].append(l)

    # Save to json
    with open(SCHEDULE_FILE, "w", encoding="utf-8") as f:
        json.dump(schedule_data, f, ensure_ascii=False, indent=2)
        
    logger.info(f"Schedule successfully saved to {SCHEDULE_FILE}")
    return schedule_data

def load_schedule() -> dict | None:
    """
    Loads parsed schedule from SCHEDULE_FILE.
    """
    if not SCHEDULE_FILE.exists():
        return None
    try:
        with open(SCHEDULE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.error(f"Failed to load schedule file: {e}")
        return None

def filter_lessons_for_week(lessons: list[dict], week_number: int, parity: str) -> list[dict]:
    """
    Filters lessons by week number and parity.
    """
    filtered = []
    for l in lessons:
        start_w = l.get("start_week", 1)
        end_w = l.get("end_week", 15)
        l_parity = l.get("parity", "all")
        
        # Check week range
        if not (start_w <= week_number <= end_w):
            continue
            
        # Check parity
        if l_parity != "all" and l_parity != parity:
            continue
            
        filtered.append(l)
    return filtered
